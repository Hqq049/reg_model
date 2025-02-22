import pandas as pd
import numpy as np
from typing import Dict, List
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side, Protection
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from pathlib import Path

class ExcelGenerator:
    def __init__(self):
        self.columns = [
            'RegisterName',
            'Offset',
            'Filed',
            'Width',
            'RW-op',
            'Reset_value',
            'Security',
            'Description'
        ]
        
        # 列宽设置
        self.column_widths = {
            'RegisterName': 15,
            'Offset': 10,
            'Filed': 15,
            'Width': 8,
            'RW-op': 12,  # 增加宽度以适应更长的访问类型
            'Reset_value': 12,
            'Security': 8,
            'Description': 40
        }
        
        # 访问类型定义
        self.access_types = {
            'RW': '读写',
            'RO': '只读',
            'WO': '只写',
            'W1C': '写1清零',
            'W1S': '写1置位',
            'RC': '读清零',
            'RS': '读置位',
            'WC': '写清零',
            'WS': '写置位',
            'W1T': '写1触发',
            'W0C': '写0清零',
            'W0S': '写0置位'
        }
        
        # 安全属性定义
        self.security_types = {
            'S': '安全访问',
            'NS': '非安全访问'
        }
        
        # 有效值范围
        self.valid_values = {
            'RW-op': list(self.access_types.keys()),
            'Security': list(self.security_types.keys())
        }
        
    def add_help_sheet(self, wb: openpyxl.Workbook):
        """添加帮助说明页"""
        help_sheet = wb.create_sheet("Help")
        
        # 设置列宽
        help_sheet.column_dimensions['A'].width = 15
        help_sheet.column_dimensions['B'].width = 50
        
        # 添加访问类型说明
        help_sheet['A1'] = "访问类型说明"
        help_sheet['A1'].font = Font(bold=True)
        row = 2
        for access_type, desc in self.access_types.items():
            help_sheet[f'A{row}'] = access_type
            help_sheet[f'B{row}'] = desc
            row += 1
            
        # 添加安全属性说明
        row += 1
        help_sheet[f'A{row}'] = "安全属性说明"
        help_sheet[f'A{row}'].font = Font(bold=True)
        row += 1
        for security_type, desc in self.security_types.items():
            help_sheet[f'A{row}'] = security_type
            help_sheet[f'B{row}'] = desc
            row += 1
            
        # 添加示例说明
        row += 2
        help_sheet[f'A{row}'] = "示例说明"
        help_sheet[f'A{row}'].font = Font(bold=True)
        row += 1
        help_sheet[f'A{row}'] = "CTRL.enable"
        help_sheet[f'B{row}'] = "RW类型，安全访问(S)，表示可读写的控制位"
        row += 1
        help_sheet[f'A{row}'] = "STATUS.irq"
        help_sheet[f'B{row}'] = "W1C类型，表示写1清零的中断状态位"
        
    def format_excel(self, file_path: Path):
        """格式化Excel文件并添加数据验证"""
        wb = openpyxl.load_workbook(str(file_path))
        ws = wb.active
        
        # 设置表头样式
        header_fill = PatternFill(start_color='CCCCCC', end_color='CCCCCC', fill_type='solid')
        header_font = Font(bold=True)
        header_align = Alignment(horizontal='center', vertical='center')
        border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
        
        # 应用表头样式
        for col in range(1, len(self.columns) + 1):
            cell = ws.cell(row=1, column=col)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = header_align
            cell.border = border
            
        # 设置列宽
        for col, name in enumerate(self.columns, 1):
            ws.column_dimensions[get_column_letter(col)].width = self.column_widths[name]
            
        # 添加数据验证
        rw_op_dv = DataValidation(
            type="list",
            formula1=f'"{",".join(self.valid_values["RW-op"])}"',
            allow_blank=False
        )
        security_dv = DataValidation(
            type="list",
            formula1=f'"{",".join(self.valid_values["Security"])}"',
            allow_blank=False
        )
        
        # 应用数据验证
        ws.add_data_validation(rw_op_dv)
        ws.add_data_validation(security_dv)
        
        # 设置数据验证范围
        rw_op_col = self.columns.index('RW-op') + 1
        security_col = self.columns.index('Security') + 1
        rw_op_dv.add(f'{get_column_letter(rw_op_col)}2:{get_column_letter(rw_op_col)}1048576')
        security_dv.add(f'{get_column_letter(security_col)}2:{get_column_letter(security_col)}1048576')
        
        # 添加帮助说明页
        self.add_help_sheet(wb)
        
        # 保存文件
        wb.save(str(file_path))
        
    def create_template(self, output_path: Path):
        """创建示例Excel模板"""
        # 创建示例数据
        data = {
            'RegisterName': [
                'CTRL',     'CTRL',     'CTRL',      # 控制寄存器
                'STATUS',   'STATUS',   'STATUS',    # 状态寄存器
                'DATA',     'DATA',                  # 数据寄存器
                'VERSION',  'VERSION'                # 版本寄存器
            ],
            'Offset': [
                '0x00',    '0x00',    '0x00',      # CTRL偏移
                '0x04',    '0x04',    '0x04',      # STATUS偏移
                '0x08',    '0x08',                 # DATA偏移
                '0x0C',    '0x0C'                  # VERSION偏移
            ],
            'Filed': [
                'enable',   'mode',     'irq_en',    # CTRL字段
                'busy',     'irq',      'ready',     # STATUS字段
                'data_h',   'data_l',               # DATA字段
                'major',    'minor'                 # VERSION字段
            ],
            'Width': [
                1,          2,          1,           # CTRL位宽
                1,          1,          1,           # STATUS位宽
                16,         16,                     # DATA位宽
                4,          4                       # VERSION位宽
            ],
            'RW-op': [
                'RW',       'RW',       'RW',        # CTRL访问权限
                'RO',       'W1C',      'RO',        # STATUS访问权限
                'RW',       'RW',                    # DATA访问权限
                'RO',       'RO'                     # VERSION访问权限
            ],
            'Reset_value': [
                0,          0,          0,           # CTRL复位值
                0,          0,          1,           # STATUS复位值
                0,          0,                      # DATA复位值
                1,          0                       # VERSION复位值
            ],
            'Security': [
                'S',        'S',        'S',         # CTRL安全属性
                'S',        'S',        'S',         # STATUS安全属性
                'NS',       'NS',                    # DATA安全属性
                'S',        'S'                      # VERSION安全属性
            ],
            'Description': [
                '模块使能控制',  '工作模式选择',  '中断使能控制',    # CTRL描述
                '忙状态指示',    '中断状态',      '就绪状态',       # STATUS描述
                '数据高16位',    '数据低16位',                    # DATA描述
                '主版本号',      '次版本号'                      # VERSION描述
            ]
        }
        
        # 创建DataFrame
        df = pd.DataFrame(data)
        
        # 保存到Excel文件
        df.to_excel(str(output_path), index=False)
        
        # 格式化Excel文件
        self.format_excel(output_path)
        
    def validate_excel(self, df: pd.DataFrame) -> List[str]:
        """验证Excel数据有效性"""
        errors = []
        
        # 检查必填字段
        for col in self.columns:
            if df[col].isnull().any():
                errors.append(f"列 {col} 存在空值")
                
        # 检查位宽有效性
        if (df['Width'] <= 0).any():
            errors.append("位宽必须大于0")
            
        # 检查读写属性有效性
        invalid_rw = df[~df['RW-op'].isin(self.valid_values['RW-op'])]
        if not invalid_rw.empty:
            errors.append(f"无效的读写属性: {invalid_rw['RW-op'].unique()}")
            
        # 检查安全属性有效性
        invalid_security = df[~df['Security'].isin(self.valid_values['Security'])]
        if not invalid_security.empty:
            errors.append("无效的安全属性值")
            
        # 检查寄存器总位宽
        reg_groups = df.groupby('RegisterName')
        for name, group in reg_groups:
            total_width = group['Width'].sum()
            if total_width > 32:
                errors.append(f"寄存器 {name} 总位宽 {total_width} 超过32位")
                
        return errors
        
    def read_excel(self, file_path: Path) -> pd.DataFrame:
        """读取并验证Excel文件"""
        df = pd.read_excel(str(file_path), usecols=self.columns)
        
        # 验证数据有效性
        errors = self.validate_excel(df)
        if errors:
            raise ValueError("\n".join(errors))
            
        return df 