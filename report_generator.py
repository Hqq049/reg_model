import pandas as pd
from typing import List, Dict
import datetime
import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from reg_model_generator import RegModelGenerator, AccessType, SecurityType, RegModel, RegField
from pathlib import Path

class ReportGenerator:
    def __init__(self, excel_file: Path):
        self.reg_model_gen = RegModelGenerator(excel_file)
        self.excel_file = excel_file
        self.report_time = datetime.datetime.now()
        self.doc = Document()
        # 设置中文字体
        self.doc.styles['Normal'].font.name = '宋体'
        self.doc.styles['Normal']._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
        
    def add_heading(self, text: str, level: int = 1):
        """添加标题"""
        heading = self.doc.add_heading(text, level=level)
        # 设置标题字体
        heading.style.font.name = '黑体'
        heading.style.element.rPr.rFonts.set(qn('w:eastAsia'), '黑体')
        
    def add_table(self, headers: List[str], rows: List[List[str]]):
        """添加表格"""
        table = self.doc.add_table(rows=1, cols=len(headers), style='Table Grid')
        
        # 添加表头
        header_cells = table.rows[0].cells
        for i, header in enumerate(headers):
            header_cells[i].text = header
            # 设置表头格式
            paragraph = header_cells[i].paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = paragraph.runs[0]
            run.font.bold = True
            
        # 添加数据行
        for row_data in rows:
            row_cells = table.add_row().cells
            for i, cell_data in enumerate(row_data):
                row_cells[i].text = str(cell_data)
                row_cells[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
                
    def generate_reg_summary(self):
        """生成寄存器概要信息"""
        self.add_heading("寄存器概要", 1)
        self.add_heading("寄存器列表", 2)
        
        # 创建寄存器列表表格
        headers = ['寄存器名称', '地址偏移', '总位宽', '字段数量', '描述']
        rows = []
        
        # 遍历字典而不是使用groupby
        for reg_name, reg_model in self.reg_model_gen.reg_models.items():
            rows.append([
                reg_name,
                f"0x{reg_model.offset:X}",
                str(reg_model.width),
                str(len(reg_model.fields)),
                reg_model.description
            ])
        
        self.add_table(headers, rows)
        self.doc.add_paragraph()
        
    def generate_field_details(self):
        """生成字段详细信息"""
        self.add_heading("寄存器详细信息", 1)
        
        headers = ['字段名称', '位宽', '偏移', '访问类型', '复位值', '安全属性', '描述']
        
        # 直接遍历字典而不是使用groupby
        for reg_name, reg_model in self.reg_model_gen.reg_models.items():
            self.add_heading(f"{reg_name} 寄存器", 2)
            
            rows = []
            for field_name, field in reg_model.fields.items():
                rows.append([
                    field_name,
                    str(field.width),
                    str(field.offset),
                    field.access_type,
                    f"0x{field.reset_value:X}",
                    '非安全' if field.security == 'NS' else '安全',
                    field.description
                ])
                
            self.add_table(headers, rows)
            self.doc.add_paragraph()
            
    def generate_coverage_summary(self):
        """生成覆盖率信息概要"""
        self.add_heading("覆盖率收集计划", 1)
        
        # 字段覆盖率
        self.add_heading("字段覆盖率", 2)
        headers = ['寄存器', '字段', '覆盖类型']
        rows = []
        
        # 直接遍历字典而不是使用groupby
        for reg_name, reg_model in self.reg_model_gen.reg_models.items():
            for field_name, field in reg_model.fields.items():
                if field.access_type == 'RW':
                    rows.append([
                        reg_name,
                        field_name,
                        '值覆盖, 转换覆盖'
                    ])
                    
        self.add_table(headers, rows)
        self.doc.add_paragraph()
        
        # 交叉覆盖
        self.add_heading("交叉覆盖", 2)
        headers = ['寄存器', '交叉字段']
        rows = []
        
        # 直接遍历字典而不是使用groupby
        for reg_name, reg_model in self.reg_model_gen.reg_models.items():
            rw_fields = [name for name, field in reg_model.fields.items() 
                        if field.access_type == 'RW']
            if len(rw_fields) > 1:
                rows.append([reg_name, ' × '.join(rw_fields)])
                
        self.add_table(headers, rows)
        self.doc.add_paragraph()
        
    def generate_file_list(self):
        """生成生成文件列表"""
        self.add_heading("生成文件列表", 1)
        headers = ['文件名', '用途']
        rows = [
            ['reg_model.sv', '寄存器模型定义'],
            ['reg_coverage.sv', '覆盖率收集代码'],
            ['reg_tests.sv', '基本测试用例'],
            ['fault_tests.sv', '错误注入测试']
        ]
        self.add_table(headers, rows)
        
    def generate_markdown(self, output_path: str):
        """生成Markdown格式的寄存器文档"""
        if not self.reg_model_gen.reg_models:
            self.reg_model_gen.load_excel()
            
        md_content = [
            "# 寄存器模型文档",
            f"\n生成时间: {datetime.datetime.now().strftime('%Y-%m-%d %H:%M:%S')}",
            f"源文件: {os.path.basename(self.excel_file)}",
            "\n## 寄存器访问类型说明",
            "| 类型 | 描述 | 硬件访问 | 软件访问 |",
            "|------|------|----------|----------|",
            "| RW   | 读写 | 读写     | 读写     |",
            "| RO   | 只读 | 只写     | 只读     |",
            "| WO   | 只写 | 只读     | 只写     |",
            "| W1C  | 写1清零 | 置位   | 写1清零 |",
            "| W1S  | 写1置位 | 清零   | 写1置位 |",
            "| RC   | 读清零 | 置位    | 读清零  |",
            "| RS   | 读置位 | 清零    | 读置位  |",
            "| WC   | 写清零 | 置位    | 写清零  |",
            "| WS   | 写置位 | 清零    | 写置位  |",
            "| W1T  | 写1触发 | -      | 写1触发 |",
            "| W0C  | 写0清零 | 置位   | 写0清零 |",
            "| W0S  | 写0置位 | 清零   | 写0置位 |",
            "\n## 安全属性说明",
            "| 类型 | 描述 |",
            "|------|------|",
            "| S    | 安全访问，只能在安全模式下访问 |",
            "| NS   | 非安全访问，可以在任何模式下访问 |",
            "\n## 寄存器描述"
        ]
        
        # 生成每个寄存器的描述
        addr_offset = 0
        for reg_name, reg_model in self.reg_model_gen.reg_models.items():
            md_content.extend([
                f"\n### {reg_name} (偏移: 0x{addr_offset:03X})",
                f"安全属性: {'安全' if reg_model.is_secure() else '非安全'}",
                f"\n总位宽: {reg_model.width} 位",
                f"复位值: 0x{reg_model.get_reset_value():X}",
                "\n| 位域 | 位宽 | 偏移 | 访问类型 | 安全属性 | 复位值 | 描述 |",
                "|------|------|------|----------|----------|--------|------|"
            ])
            
            # 添加每个字段的描述
            for field_name, field in reg_model.fields.items():
                md_content.append(
                    f"| {field_name} | {field.width} | {field.offset} | "
                    f"{field.access_type} | {field.security} | "
                    f"0x{field.reset_value:X} | {field.description} |"
                )
                
            addr_offset += 4
            
        # 写入文件
        with open(output_path, 'w', encoding='utf-8') as f:
            f.write('\n'.join(md_content))
            
    def generate_html(self, output_path: str):
        """生成HTML格式的寄存器文档"""
        if not self.reg_model_gen.reg_models:
            self.reg_model_gen.load_excel()
            
        html_content = ["""
<!DOCTYPE html>
<html>
<head>
    <meta charset="utf-8">
    <title>寄存器模型文档</title>
    <style>
        body { font-family: Arial, sans-serif; margin: 20px; }
        table { border-collapse: collapse; width: 100%; margin: 10px 0; }
        th, td { border: 1px solid #ddd; padding: 8px; text-align: left; }
        th { background-color: #f2f2f2; }
        h1 { color: #333; }
        h2 { color: #666; margin-top: 20px; }
        h3 { color: #999; }
        .reg-info { margin: 10px 0; }
        .field-table { margin-left: 20px; }
    </style>
</head>
<body>
    <h1>寄存器模型文档</h1>"""]
        
        html_content.extend([
            f"    <p>生成时间: {datetime.datetime.now().strftime('%Y-%m-%d %H:%M:%S')}</p>",
            f"    <p>源文件: {os.path.basename(self.excel_file)}</p>",
            "    <h2>寄存器访问类型说明</h2>",
            "    <table>",
            "        <tr><th>类型</th><th>描述</th><th>硬件访问</th><th>软件访问</th></tr>"
        ])
        
        # 添加访问类型说明
        access_types = {
            "RW": ["读写", "读写", "读写"],
            "RO": ["只读", "只写", "只读"],
            "WO": ["只写", "只读", "只写"],
            "W1C": ["写1清零", "置位", "写1清零"],
            "W1S": ["写1置位", "清零", "写1置位"],
            "RC": ["读清零", "置位", "读清零"],
            "RS": ["读置位", "清零", "读置位"],
            "WC": ["写清零", "置位", "写清零"],
            "WS": ["写置位", "清零", "写置位"],
            "W1T": ["写1触发", "-", "写1触发"],
            "W0C": ["写0清零", "置位", "写0清零"],
            "W0S": ["写0置位", "清零", "写0置位"]
        }
        
        for access_type, desc in access_types.items():
            html_content.append(
                f"        <tr><td>{access_type}</td><td>{desc[0]}</td>"
                f"<td>{desc[1]}</td><td>{desc[2]}</td></tr>"
            )
            
        html_content.extend([
            "    </table>",
            "    <h2>安全属性说明</h2>",
            "    <table>",
            "        <tr><th>类型</th><th>描述</th></tr>",
            "        <tr><td>S</td><td>安全访问，只能在安全模式下访问</td></tr>",
            "        <tr><td>NS</td><td>非安全访问，可以在任何模式下访问</td></tr>",
            "    </table>",
            "    <h2>寄存器描述</h2>"
        ])
        
        # 生成每个寄存器的描述
        addr_offset = 0
        for reg_name, reg_model in self.reg_model_gen.reg_models.items():
            html_content.extend([
                f"    <h3>{reg_name} (偏移: 0x{addr_offset:03X})</h3>",
                "    <div class='reg-info'>",
                f"        <p>安全属性: {'安全' if reg_model.is_secure() else '非安全'}</p>",
                f"        <p>总位宽: {reg_model.width} 位</p>",
                f"        <p>复位值: 0x{reg_model.get_reset_value():X}</p>",
                "    </div>",
                "    <table class='field-table'>",
                "        <tr><th>位域</th><th>位宽</th><th>偏移</th><th>访问类型</th>"
                "<th>安全属性</th><th>复位值</th><th>描述</th></tr>"
            ])
            
            # 添加每个字段的描述
            for field_name, field in reg_model.fields.items():
                html_content.append(
                    f"        <tr><td>{field_name}</td><td>{field.width}</td>"
                    f"<td>{field.offset}</td><td>{field.access_type}</td>"
                    f"<td>{field.security}</td><td>0x{field.reset_value:X}</td>"
                    f"<td>{field.description}</td></tr>"
                )
                
            html_content.append("    </table>")
            addr_offset += 4
            
        html_content.append("</body>\n</html>")
        
        # 写入文件
        with open(output_path, 'w', encoding='utf-8') as f:
            f.write('\n'.join(html_content))
            
    def generate_word(self, output_path: str):
        """生成Word格式的寄存器文档"""
        if not self.reg_model_gen.reg_models:
            self.reg_model_gen.load_excel()
            
        # 添加文档标题
        self.add_heading("寄存器模型文档", 0)
        self.doc.add_paragraph(f"生成时间: {self.report_time.strftime('%Y-%m-%d %H:%M:%S')}")
        self.doc.add_paragraph(f"源文件: {os.path.basename(self.excel_file)}")
        
        # 添加访问类型说明
        self.add_heading("寄存器访问类型说明", 1)
        headers = ['类型', '描述', '硬件访问', '软件访问']
        rows = [
            ['RW', '读写', '读写', '读写'],
            ['RO', '只读', '只写', '只读'],
            ['WO', '只写', '只读', '只写'],
            ['W1C', '写1清零', '置位', '写1清零'],
            ['W1S', '写1置位', '清零', '写1置位'],
            ['RC', '读清零', '置位', '读清零'],
            ['RS', '读置位', '清零', '读置位'],
            ['WC', '写清零', '置位', '写清零'],
            ['WS', '写置位', '清零', '写置位'],
            ['W1T', '写1触发', '-', '写1触发'],
            ['W0C', '写0清零', '置位', '写0清零'],
            ['W0S', '写0置位', '清零', '写0置位']
        ]
        self.add_table(headers, rows)
        self.doc.add_paragraph()
        
        # 添加安全属性说明
        self.add_heading("安全属性说明", 1)
        headers = ['类型', '描述']
        rows = [
            ['S', '安全访问，只能在安全模式下访问'],
            ['NS', '非安全访问，可以在任何模式下访问']
        ]
        self.add_table(headers, rows)
        self.doc.add_paragraph()
        
        # 添加寄存器描述
        self.add_heading("寄存器描述", 1)
        addr_offset = 0
        for reg_name, reg_model in self.reg_model_gen.reg_models.items():
            # 添加寄存器标题
            self.add_heading(f"{reg_name} (偏移: 0x{addr_offset:03X})", 2)
            
            # 添加寄存器基本信息
            info_para = self.doc.add_paragraph()
            info_para.add_run(f"安全属性: {'安全' if reg_model.is_secure() else '非安全'}\n")
            info_para.add_run(f"总位宽: {reg_model.width} 位\n")
            info_para.add_run(f"复位值: 0x{reg_model.get_reset_value():X}")
            
            # 添加字段表格
            headers = ['位域', '位宽', '偏移', '访问类型', '安全属性', '复位值', '描述']
            rows = []
            for field_name, field in reg_model.fields.items():
                rows.append([
                    field_name,
                    str(field.width),
                    str(field.offset),
                    field.access_type,
                    field.security,
                    f"0x{field.reset_value:X}",
                    field.description
                ])
            self.add_table(headers, rows)
            self.doc.add_paragraph()
            addr_offset += 4
            
        # 保存文档
        self.doc.save(output_path)

    def generate_all(self, output_dir: str):
        """生成所有格式的文档"""
        if not os.path.exists(output_dir):
            os.makedirs(output_dir)
            
        # 生成各种格式的文档
        self.generate_markdown(f"{output_dir}/register_doc.md")
        self.generate_html(f"{output_dir}/register_doc.html")
        self.generate_word(f"{output_dir}/register_doc.docx")
        
        # 生成完整的技术报告
        self.generate_report(f"{output_dir}/technical_report.docx")
        
    def generate_report(self, output_path: str):
        """生成完整的报告"""
        if self.reg_model_gen.reg_models is None:
            self.reg_model_gen.load_excel()
            
        # 添加报告标题
        self.add_heading("寄存器模型生成报告", 0)
        self.doc.add_paragraph(f"生成时间: {self.report_time.strftime('%Y-%m-%d %H:%M:%S')}")
        self.doc.add_paragraph()
        
        # 生成各个部分
        self.generate_reg_summary()
        self.generate_field_details()
        self.generate_coverage_summary()
        self.generate_file_list()
        
        # 保存文档
        self.doc.save(output_path) 